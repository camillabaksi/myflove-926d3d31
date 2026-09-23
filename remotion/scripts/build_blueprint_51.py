"""Build Living Blueprint 5.1 from the established 5.0 document."""
from pathlib import Path
import subprocess
from docx import Document
from docx.enum.text import WD_BREAK

ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "public/blueprint/MyFlove_Living_Blueprint_5.0.docx"
OUT = Path("/mnt/documents/MyFlove_Living_Blueprint_5.1.docx")

SECTIONS = [
    ("Tillägg till version 5.1", [
        "Detta tillägg beskriver den nuvarande produkten per 23 september 2026. Blueprint 5.0:s grundprinciper består. Version 5.1 dokumenterar den nya informationsarkitekturen och de funktioner som har byggts efter 5.0.",
        "Kärnan är oförändrad: MyFlove ska hjälpa användaren att se sin faktiska kapacitet, förstå dagens helhet och hitta nästa rimliga steg utan skuld.",
    ]),
    ("Informationsarkitektur: fem tydliga vägar", [
        "Primär navigation är Översikt, Flow, Hälsa, Återhämtning och Mer. QuickFlow ligger inne i Flow. Studier, Jobb, Familj, Hem och andra områden samlas under Mer i stället för att konkurrera i huvudnavigationen.",
        "VISA → KLICKA → FÖRDJUPA är fortsatt vägledande: det viktigaste syns först, nästa handling är tydlig och detaljer öppnas först när användaren ber om dem.",
    ]),
    ("Ett gemensamt kapacitets- och energisystem", [
        "Mitt batteri, mående, FlowScore, Min dag, kalendern och tidigare frivilliga energiloggar används tillsammans. Ingen separat energimodell skapas för enskilda livsområden.",
        "Efter en genomförd aktivitet kan användaren frivilligt ange om den gav energi, tog energi, både och eller är oklar. Samma aktivitet loggas bara en gång. Insikter använder användarens egna loggar och formulerar sig försiktigt; stress och konflikt får aldrig beskrivas som energikällor.",
    ]),
    ("Flow: anpassning med användaren i kontroll", [
        "Flow föreslår ett nästa rimligt steg utifrån aktuell kapacitet, dagens aktiviteter och fasta tider. Förslag kan visas, ändras, nekas eller skjutas upp. Fasta aktiviteter flyttas eller bokas inte om automatiskt.",
        "När dagen är tät kan Flow leta efter realistiska luckor för återhämtning och ta hänsyn till förberedelse, resa och övergångstid. Inget bokas utan användarens godkännande.",
    ]),
    ("Återhämtning", [
        "Återhämtning är en huvuddel med Akut Lugn, Mikropaus, Ljusboost och Stopp. Ljusboost är en sensorisk paus med varm ljuspuls, andningsrytm, valfri havsmiljö och tidsval — inte medicinsk ljusterapi.",
        "Efter en paus kan användaren göra samma frivilliga energi- och måendecheck som efter andra aktiviteter. Användningen av de olika pausvägarna kan följas anonymt i administrationen.",
    ]),
    ("Insikter", [
        "Insikter visar energi- och måendetrend för vecka och månad, upp till åtta veckor och sex månader. Perioder utan data lämnas tomma.",
        "Språket är beskrivande och icke-diagnostiskt. MyFlove ska aldrig hitta på ett historiskt mönster eller påstå att en aktivitet brukar ge energi när underlaget inte stödjer det.",
    ]),
    ("Family Flow och visuellt barnschema", [
        "Family Flow återanvänder befintliga barnprofiler, aktiviteter, kalendern, Min dag, QuickFlow, Flow och Nästa rimliga steg. Barnprofilen innehåller namn, valfri avatar och privata observationer om vad som brukar hjälpa — inga diagnos- eller medicinfält.",
        "Det visuella barnschemat visar vecka, idag, Nu och Nästa med pictogram, områdesfärger, aktivitetsbank, egna privata bilder och skyddad veckokopiering utan dubbletter. Aktuell tid räknas från en gemensam lokal tidskälla och uppdateras dynamiskt.",
        "Barnaktiviteter får aldrig automatiska energikostnader. Användaren kan frivilligt markera föräldranärvaro, förberedelse, resa, övergångstid och återhämtning efteråt. Endast dessa faktiska föräldraåtaganden vägs mot förälderns befintliga batteri.",
    ]),
    ("Träning och studier", [
        "Träningsprogram samlar Skapa mitt pass, sparade pass och historik. Pass kan anpassas efter tid, utrustning, fokus, begränsningar och aktuell kapacitet. Ett automatiskt följ-läge guidar genom övning, vila och nästa steg. Övningsbiblioteket har realistiska instruktionsbilder.",
        "Studier har redigerbart schema även efter filimport. Användaren kan välja flashcards, quiz eller båda. Originalfil, OCR-resultat, modeller och befintliga vägar bevaras.",
    ]),
    ("Språk, integritet och tillförlitlighet", [
        "Appen kan växla mellan svenska och engelska från inloggningen och Inställningar. Översättning använder ordbok först, cache och säker återgång till svenska när en översättning saknas.",
        "Familjedelning är privat som standard och sker sak för sak. Pitch- och investerarmaterial ligger i en privat adminsektion med tillfälliga länkar och återanvänd befintlig adminbehörighet.",
        "Molnsynk slår ihop aktivitetslistor med stabila objekt-id:n i stället för att låta en äldre lokal kopia skriva över nyare data. Tidigare versioner kan bevaras som privata revisioner.",
    ]),
    ("Produkt- och affärsstatus", [
        "MyFlove är en fungerande produkt. Marknad, betalningsvilja, retention och effekt ska fortsatt valideras innan de presenteras som bevisade resultat.",
        "B2C Free/Premium är utgångspunkten. Arbetsgivare och offentlig sektor är möjliga framtida spår. Ett arbetscase om 250 000 euro för 18 månaders validering används i pitchmaterialet och ska alltid märkas som ett antagande, inte som ett verifierat utfall.",
        "Verifierade externa marknadsuppskattningar som återanvänds är över 140 miljoner vuxna globalt med ADHD och omkring 3 procent av vuxna i Sverige. Källorna ska anges och siffrorna får inte presenteras som MyFloves traction.",
    ]),
]

def build():
    doc = Document(SRC)
    # Keep the established blueprint design while making the new file's cover truthful.
    for paragraph in doc.paragraphs[:12]:
        if "Version 5.0" in paragraph.text:
            for run in paragraph.runs:
                run.text = run.text.replace("Version 5.0", "Version 5.1")
        if "Om version 5.0" in paragraph.text:
            for run in paragraph.runs:
                run.text = run.text.replace("Om version 5.0", "Om version 5.1")
    doc.add_page_break()
    for index, (title, paragraphs) in enumerate(SECTIONS):
        doc.add_heading(title, level=1 if index == 0 else 2)
        for text in paragraphs:
            doc.add_paragraph(text)
    doc.add_heading("Versionsnotering", level=2)
    doc.add_paragraph("5.1 · Ny informationsarkitektur, gemensam energiloop, Återhämtning, Insikter, Family Flow, visuellt barnschema, träningsläge, redigerbart studieschema, språkstöd och stärkt synk.")
    doc.save(OUT)
    subprocess.run(["python3", "/tmp/run_libreoffice.py", "--headless", "--convert-to", "pdf", "--outdir", "/mnt/documents", str(OUT)], check=True)
    print(OUT)

if __name__ == "__main__":
    build()