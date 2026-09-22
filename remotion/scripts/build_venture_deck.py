"""Bygger en redigerbar Venture Cup-pitchdeck (PPTX) med samma innehåll som pitchfilmerna."""
from pathlib import Path
import importlib.util

from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parents[1]
LIVE = ROOT / "public/live"

spec = importlib.util.spec_from_file_location("vp", Path(__file__).with_name("build_venture_pitch.py"))
vp = importlib.util.module_from_spec(spec)
spec.loader.exec_module(vp)

INK = RGBColor(0x1F, 0x11, 0x4E)
VIOLET = RGBColor(0x77, 0x52, 0xB8)
PINK = RGBColor(0xEF, 0xBE, 0xDC)
LILAC = RGBColor(0xDD, 0xCF, 0xF4)
PANEL = RGBColor(0xF1, 0xEA, 0xF9)
MUTED = RGBColor(0x67, 0x5E, 0x75)
BG = RGBColor(0xF8, 0xF6, 0xFB)
PAPER = RGBColor(0xFF, 0xFF, 0xFF)

W = Emu(12192000)  # 16:9 13.333in
H = Emu(6858000)


def inch(v):
    return Emu(int(v * 914400))


def solid(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False


def textbox(slide, x, y, w, h, text, size, color, bold=False, font="Inter", italic=False, spacing=1.15):
    tb = slide.shapes.add_textbox(inch(x), inch(y), inch(w), inch(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = spacing
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.name = font
        r.font.color.rgb = color
    return tb


def build(lang):
    data = vp.SV if lang == "SV" else vp.EN
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    blank = prs.slide_layouts[6]

    for idx, (label, title, body, shot, note) in enumerate(data):
        s = prs.slides.add_slide(blank)
        bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
        solid(bg, BG)

        blob = s.shapes.add_shape(MSO_SHAPE.OVAL, inch(9.6), inch(-1.9), inch(4.8), inch(3.4))
        solid(blob, PINK)
        blob2 = s.shapes.add_shape(MSO_SHAPE.OVAL, inch(-2.2), inch(5.1), inch(5.2), inch(3.4))
        solid(blob2, LILAC)

        pill = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, inch(0.85), inch(0.55), inch(0.14 * len(label) + 0.5), inch(0.36))
        solid(pill, INK)
        tf = pill.text_frame
        tf.word_wrap = False
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = label
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.name = "Inter"
        r.font.color.rgb = PAPER

        textbox(s, 0.85, 1.15, 6.4, 1.9, title.replace("|", "\n"), 40, INK, bold=True, font="Plus Jakarta Sans", spacing=1.05)
        textbox(s, 0.85, 3.05, 6.4, 1.6, body, 16, MUTED, font="Inter", spacing=1.3)
        textbox(s, 0.85, 4.7, 6.4, 1.1, note.replace("\n", "  "), 17, VIOLET, italic=True, font="Caveat", spacing=1.15)

        panel = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, inch(8.1), inch(0.7), inch(4.4), inch(5.9))
        solid(panel, PANEL)
        img = LIVE / shot
        if img.exists():
            ph = 5.2
            pw = ph * 430 / 932
            frame = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, inch(8.1 + (4.4 - pw) / 2 - 0.08), inch(1.05), inch(pw + 0.16), inch(ph + 0.16))
            solid(frame, INK)
            s.shapes.add_picture(str(img), inch(8.1 + (4.4 - pw) / 2), inch(1.13), inch(pw), inch(ph))

        textbox(s, 0.85, 6.25, 3.0, 0.4, "MyFlove", 14, INK, bold=True, font="Plus Jakarta Sans")
        foot = textbox(s, 9.3, 6.3, 3.2, 0.4, "VENTURE CUP 2026", 10, VIOLET, bold=True, font="Inter")
        foot.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT

        s.notes_slide.notes_text_frame.text = f"{vp.DURATIONS[idx]} sekunder i filmversionen."

    out = Path("/mnt/documents") / f"MyFlove_Venture_Cup_2026_Redigerbar_{lang}_16x9_v2.pptx"
    prs.save(out)
    print(out)


if __name__ == "__main__":
    build("SV")
    build("EN")
