"""Build bilingual editable pitch scripts and package index."""
from pathlib import Path
import subprocess
from docx import Document
from pitch_content_51 import SV, EN

OUT = Path("/mnt/documents")

def script_doc(lang, data):
    labels = {
        "SV": ("MyFlove Pitchmanus 5.1", "Redigerbart manus för pitchdeck och treminutersfilm", "Bild", "Talmanus / huvudbudskap", "Anteckning"),
        "EN": ("MyFlove Pitch Script 5.1", "Editable script for the pitch deck and three-minute video", "Scene", "Narration / key message", "Note"),
    }[lang]
    doc = Document()
    doc.add_heading(labels[0], 0)
    doc.add_paragraph(labels[1])
    for i, (eyebrow, title, body, _shot, note) in enumerate(data, 1):
        doc.add_heading(f"{labels[2]} {i}: {title.replace('|', ' ')}", level=1)
        doc.add_paragraph(eyebrow)
        doc.add_heading(labels[3], level=2)
        doc.add_paragraph(body)
        doc.add_heading(labels[4], level=2)
        doc.add_paragraph(note)
    path = OUT / f"MyFlove_Pitchmanus_5.1_{lang}.docx"
    doc.save(path)
    subprocess.run(["python3", "/tmp/run_libreoffice.py", "--headless", "--convert-to", "pdf", "--outdir", str(OUT), str(path)], check=True)
    return path

def package_index():
    text = """# MyFlove huvudpaket 5.1

Aktuellt per 23 september 2026. Äldre versioner är historik och har inte skrivits över.

## Pitchdeck
- MyFlove_Pitchdeck_5.1_SV_16x9.pptx / .pdf — redigerbart svenskt huvuddeck
- MyFlove_Pitchdeck_5.1_EN_16x9.pptx / .pdf — redigerbart engelskt huvuddeck

## Pitchfilmer
- MyFlove_Pitchfilm_5.1_SV_16x9.mp4
- MyFlove_Pitchfilm_5.1_SV_9x16.mp4
- MyFlove_Pitchfilm_5.1_EN_16x9.mp4
- MyFlove_Pitchfilm_5.1_EN_9x16.mp4

Filmerna är cirka tre minuter, utan ljud, med mjuka övergångar och verkliga MyFlove-vyer.

## Blueprint
- MyFlove_Living_Blueprint_5.1.docx / .pdf — Blueprint 5.0 med ett nytt, samlat 5.1-tillägg

## Stöddokument
- MyFlove_Pitchmanus_5.1_SV.docx / .pdf
- MyFlove_Pitchmanus_5.1_EN.docx / .pdf

## Sifferprincip
Endast tidigare belagda externa uppskattningar återanvänds: över 140 miljoner vuxna globalt och omkring 3 procent av vuxna i Sverige uppskattas leva med ADHD. Kapitalbehovet 250 000 euro och 18 månaders runway presenteras uttryckligen som arbetsantaganden. Inga användar-, intäkts- eller effektsiffror presenteras som bevisade innan de har mätts.
"""
    path = OUT / "MyFlove_Huvudpaket_5.1_INNEHALL.md"
    path.write_text(text)
    return path

if __name__ == "__main__":
    print(script_doc("SV", SV))
    print(script_doc("EN", EN))
    print(package_index())