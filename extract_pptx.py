import sys
from pptx import Presentation

def extract_text(file_path):
    print(f"--- Content of {file_path} ---")
    prs = Presentation(file_path)
    for i, slide in enumerate(prs.slides):
        print(f"Slide {i+1}:")
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text = shape.text.strip()
                if text:
                    print(f"  [Text]: {text}")
    print("\n")

if __name__ == "__main__":
    extract_text("public/pitch/MyFlove_Pitchdeck_5.0.pptx")
    extract_text("public/pitch/MyFlove_Investor_Pitch_3min_EN.pptx")
