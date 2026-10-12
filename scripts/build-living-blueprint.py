"""Build the current downloadable blueprint from its Markdown source."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
root = Path(__file__).resolve().parents[1]
doc = Document()
sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = Inches(.65)
sec.left_margin = sec.right_margin = Inches(.8)
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.12
for name in ['Title','Heading 1']:
    st=doc.styles[name]; st.font.name='Calibri'; st.font.color.rgb=RGBColor.from_string('512654')
doc.styles['Title'].font.size=Pt(30)
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
doc.styles['Caption'].font.color.rgb=RGBColor.from_string('512654')
doc.styles['Heading 1'].font.size=Pt(18)
sec.header.paragraphs[0].text='PACEABLE  /  LIVING BLUEPRINT'
sec.header.paragraphs[0].style='Caption'
f=sec.footer.paragraphs[0]; f.text='Version 5.1 · 12 oktober 2026                                      '
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); f._p.append(field)
for line in (root/'docs/living-blueprint.md').read_text().splitlines():
    if not line: continue
    if line.startswith('# '): doc.add_paragraph(line[2:],'Title')
    elif line.startswith('## '):
        if line.startswith(('## 3.','## 5.')): doc.add_page_break()
        doc.add_paragraph(line[3:],'Heading 1')
    else: doc.add_paragraph(line)
doc.core_properties.title='Paceable Living Blueprint 5.1'
doc.core_properties.subject='Aktuell produktbeskrivning och publiceringsstatus'
doc.core_properties.author='Paceable'
out=root/'public/blueprint/Paceable_Living_Blueprint_5.1.docx'
doc.save(out)
print(out)
